"""DOCX converter using python-docx."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.table import Table

from omnivorous.converters.base import BaseConverter
from omnivorous.models import ConvertResult, DocumentMetadata
from omnivorous.tokens import count_tokens, get_encoding_name

_HEADING_MAP = {
    "Heading 1": "#",
    "Heading 2": "##",
    "Heading 3": "###",
    "Heading 4": "####",
    "Heading 5": "#####",
    "Heading 6": "######",
}


def _table_to_markdown(table: Table) -> str:
    """Convert a docx Table to a markdown table."""
    rows: list[list[str]] = []
    for row in table.rows:
        rows.append([cell.text.strip() for cell in row.cells])
    if not rows:
        return ""

    lines: list[str] = []
    header = rows[0]
    lines.append("| " + " | ".join(header) + " |")
    lines.append("| " + " | ".join("---" for _ in header) + " |")
    for row in rows[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


def _extract_image_placeholders(element) -> list[str]:
    """Extract image placeholders from a paragraph element's w:drawing descendants."""
    from docx.oxml.ns import qn

    placeholders: list[str] = []
    for drawing in element.iter(qn("w:drawing")):
        # Look for wp:docPr which carries alt text (descr) and name
        for doc_pr in drawing.iter(qn("wp:docPr")):
            alt = doc_pr.get("descr") or doc_pr.get("name") or "image"
            placeholders.append(f"![{alt}]()")
            break
        else:
            placeholders.append("![image]()")
    return placeholders


class DocxConverter(BaseConverter):
    @property
    def name(self) -> str:
        return "DOCX"

    def convert(self, path: Path) -> ConvertResult:
        doc = Document(str(path))
        parts: list[str] = []
        headings: list[str] = []
        table_count = 0
        image_count = 0

        # Interleave paragraphs and tables in document order
        # python-docx exposes doc.element.body which contains all elements in order
        from docx.oxml.ns import qn

        for element in doc.element.body:
            if element.tag == qn("w:p"):
                # It's a paragraph
                from docx.text.paragraph import Paragraph

                para = Paragraph(element, doc)
                style_name = para.style.name if para.style else ""
                text = para.text.strip()

                # Extract image placeholders from this paragraph
                img_placeholders = _extract_image_placeholders(element)
                image_count += len(img_placeholders)

                if not text and not img_placeholders:
                    continue
                prefix = _HEADING_MAP.get(style_name, "")
                if prefix and text:
                    headings.append(f"{prefix} {text}")
                    parts.append(f"{prefix} {text}")
                elif text:
                    parts.append(text)
                for placeholder in img_placeholders:
                    parts.append(placeholder)
            elif element.tag == qn("w:tbl"):
                table = Table(element, doc)
                md = _table_to_markdown(table)
                if md:
                    table_count += 1
                    parts.append(md)

        content = "\n\n".join(parts)
        metadata = DocumentMetadata(
            source=path.name,
            format="docx",
            title=path.stem,
            headings=headings,
            tables=table_count,
            images=image_count,
            tokens_estimate=count_tokens(content),
            encoding=get_encoding_name(),
        )
        return ConvertResult(content=content, metadata=metadata)
